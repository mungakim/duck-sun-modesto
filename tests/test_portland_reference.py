"""Regression tests for the Portland, OR reference row in the Excel report.

Portland is a side reference: one row, own visual band, and explicitly NOT
part of the Modesto weighted consensus.
"""

from pathlib import Path

from openpyxl import load_workbook

from duck_sun.excel_report import generate_excel_report

PORTLAND_BANNER_ROW = 23
PORTLAND_DATA_ROW = 24
SOLAR_TITLE_ROW = 26
SOLAR_HEADER_ROW = 27


def _inputs(days=8):
    dates = [f"2026-08-{d:02d}" for d in range(1, days + 1)]
    om_data = {
        "daily_forecast": [
            {
                "date": d,
                "day_name": "Saturday",
                "high_f": 95 + i,
                "low_f": 60 + i,
                "precip_prob": 5,
                "condition": "Sunny",
            }
            for i, d in enumerate(dates)
        ],
        "hourly": [],
    }
    google_data = {
        "daily": [
            {"date": d, "high_f": 96 + i, "low_f": 61 + i, "condition": "Sunny", "precip_prob": 5}
            for i, d in enumerate(dates)
        ],
        "hourly": [],
    }
    portland_data = {
        # Portland runs much cooler - values are unmistakable in the sheet
        "daily": [
            {"date": d, "high_f": 78 + i, "low_f": 55 + i, "condition": "Cloudy", "precip_prob": 30}
            for i, d in enumerate(dates)
        ],
        "hourly": [],
    }
    return om_data, google_data, portland_data, dates


def _render(tmp_path: Path, portland_data, days=8):
    om_data, google_data, _, _ = _inputs(days)
    out = tmp_path / "report.xlsx"
    generate_excel_report(
        om_data=om_data,
        noaa_data=[],
        met_data=[],
        accu_data=[],
        google_data=google_data,
        portland_data=portland_data,
        output_path=out,
    )
    return load_workbook(out).active


def test_portland_row_renders_high_low_for_every_day_column(tmp_path: Path):
    _, _, portland_data, _ = _inputs()
    ws = _render(tmp_path, portland_data)

    # Day columns start at E (col(3) with COL_OFFSET=2); hi/lo pairs per day
    expected = []
    for i in range(8):
        expected.extend([78 + i, 55 + i])

    letters = "EFGHIJKLMNOPQRST"
    actual = [ws[f"{c}{PORTLAND_DATA_ROW}"].value for c in letters]
    assert actual == expected, f"Portland row mismatch: {actual}"


def test_portland_is_labelled_and_marked_reference_only(tmp_path: Path):
    _, _, portland_data, _ = _inputs()
    ws = _render(tmp_path, portland_data)

    banner = str(ws[f"C{PORTLAND_BANNER_ROW}"].value)
    assert "PORTLAND, OR" in banner
    assert "REFERENCE" in banner.upper()
    assert "not included in the Modesto consensus" in banner

    assert ws[f"C{PORTLAND_DATA_ROW}"].value == "PORTLAND, OR"


def test_portland_values_are_excluded_from_the_modesto_weighted_average(tmp_path: Path):
    """The consensus formula must still span only source rows 13-19."""
    _, _, portland_data, _ = _inputs()
    ws = _render(tmp_path, portland_data)

    for col_letter in ("E", "F", "G", "H"):
        formula = ws[f"{col_letter}20"].value
        assert isinstance(formula, str) and formula.startswith("=")
        assert f"{col_letter}13:{col_letter}19" in formula
        assert str(PORTLAND_DATA_ROW) not in formula.replace(f"{col_letter}13:{col_letter}19", "")


def test_missing_portland_data_blanks_the_row_without_crashing(tmp_path: Path):
    ws = _render(tmp_path, None)

    assert ws[f"C{PORTLAND_DATA_ROW}"].value == "PORTLAND, OR"
    for c in "EFGHIJKLMNOPQRST":
        assert ws[f"{c}{PORTLAND_DATA_ROW}"].value == "--"


def test_portland_gaps_render_as_placeholders_not_shifted_values(tmp_path: Path):
    """A short Portland forecast must leave later columns blank, not slide left."""
    _, _, portland_data, dates = _inputs()
    portland_data["daily"] = portland_data["daily"][:3]
    ws = _render(tmp_path, portland_data)

    assert ws[f"E{PORTLAND_DATA_ROW}"].value == 78
    assert ws[f"I{PORTLAND_DATA_ROW}"].value == 80   # day index 2 high
    assert ws[f"K{PORTLAND_DATA_ROW}"].value == "--"  # day index 3 high, no data


def test_solar_block_shifted_below_portland(tmp_path: Path):
    _, _, portland_data, _ = _inputs()
    ws = _render(tmp_path, portland_data)

    assert "SOLAR FORECAST" in str(ws[f"D{SOLAR_TITLE_ROW}"].value)
    assert ws[f"D{SOLAR_HEADER_ROW}"].value == "DATE"
    assert ws[f"E{SOLAR_HEADER_ROW}"].value == "9AM"
    # Solar data starts immediately under the header and must not collide
    # with the Portland band above it.
    assert ws[f"D{PORTLAND_DATA_ROW}"].value is None or ws[f"D{PORTLAND_DATA_ROW}"].value == ""
