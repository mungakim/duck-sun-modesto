"""Regression tests for the SIDE REFERENCE TEMPS band in the Excel report.

Portland, OR and Phoenix, AZ each get one Hi/Lo row under a shared banner and
day-name row. Both are references only: own visual band, own colours, and
explicitly NOT part of the Modesto weighted consensus.
"""

from pathlib import Path

from openpyxl import load_workbook

from duck_sun.excel_report import generate_excel_report

SIDE_REF_BANNER_ROW = 22
SIDE_REF_DAYS_ROW = 23
PORTLAND_DATA_ROW = 24
PHOENIX_DATA_ROW = 25
SOLAR_TITLE_ROW = 27
SOLAR_HEADER_ROW = 28
SOLAR_LEGEND_ROW = 44   # header + 7 days x 2 rows + 1 spacer

# Day columns start at E (col(3) with COL_OFFSET=2); hi/lo pairs per day
DAY_COLUMNS = "EFGHIJKLMNOPQRST"


def _inputs(days=8):
    dates = [f"2026-08-{d:02d}" for d in range(1, days + 1)]
    om_data = {
        "daily_forecast": [
            {
                "date": d,
                "day_name": "Saturday",
                "high_f": 95 + i,
                "low_f": 60 + i,
                "precip_prob": 5,
                "condition": "Sunny",
            }
            for i, d in enumerate(dates)
        ],
        "hourly": [],
    }
    google_data = {
        "daily": [
            {"date": d, "high_f": 96 + i, "low_f": 61 + i, "condition": "Sunny", "precip_prob": 5}
            for i, d in enumerate(dates)
        ],
        "hourly": [],
    }
    portland_data = {
        # Portland runs much cooler - values are unmistakable in the sheet
        "daily": [
            {"date": d, "high_f": 78 + i, "low_f": 55 + i, "condition": "Cloudy", "precip_prob": 30}
            for i, d in enumerate(dates)
        ],
        "hourly": [],
    }
    phoenix_data = {
        # Phoenix runs much hotter - also unmistakable, and distinct from Portland
        "daily": [
            {"date": d, "high_f": 105 + i, "low_f": 80 + i, "condition": "Sunny", "precip_prob": 0}
            for i, d in enumerate(dates)
        ],
        "hourly": [],
    }
    return om_data, google_data, portland_data, phoenix_data, dates


def _render(tmp_path: Path, portland_data, phoenix_data, days=8):
    om_data, google_data, _, _, _ = _inputs(days)
    out = tmp_path / "report.xlsx"
    generate_excel_report(
        om_data=om_data,
        noaa_data=[],
        accu_data=[],
        google_data=google_data,
        portland_data=portland_data,
        phoenix_data=phoenix_data,
        output_path=out,
    )
    return load_workbook(out).active


def _render_default(tmp_path: Path):
    _, _, portland_data, phoenix_data, _ = _inputs()
    return _render(tmp_path, portland_data, phoenix_data)


def _row_values(ws, row):
    return [ws[f"{c}{row}"].value for c in DAY_COLUMNS]


def _rgb(cell):
    return str(cell.fill.start_color.rgb)[-6:].upper()


def test_portland_row_renders_high_low_for_every_day_column(tmp_path: Path):
    ws = _render_default(tmp_path)

    expected = []
    for i in range(8):
        expected.extend([78 + i, 55 + i])
    assert _row_values(ws, PORTLAND_DATA_ROW) == expected


def test_phoenix_row_renders_high_low_for_every_day_column(tmp_path: Path):
    ws = _render_default(tmp_path)

    expected = []
    for i in range(8):
        expected.extend([105 + i, 80 + i])
    assert _row_values(ws, PHOENIX_DATA_ROW) == expected


def test_phoenix_sits_directly_below_portland_under_one_shared_header(tmp_path: Path):
    ws = _render_default(tmp_path)

    assert ws[f"C{PORTLAND_DATA_ROW}"].value == "PORTLAND, OR"
    assert ws[f"C{PHOENIX_DATA_ROW}"].value == "PHOENIX, AZ"
    assert PHOENIX_DATA_ROW == PORTLAND_DATA_ROW + 1
    assert SIDE_REF_BANNER_ROW < SIDE_REF_DAYS_ROW < PORTLAND_DATA_ROW


def test_day_names_repeat_directly_above_the_reference_rows(tmp_path: Path):
    """The Modesto header is a dozen rows up; the band needs its own day labels."""
    ws = _render_default(tmp_path)

    # Day 0 is always TODAY; the rest are 3-letter day abbreviations
    assert ws[f"E{SIDE_REF_DAYS_ROW}"].value == "TODAY"
    for col_letter in ("G", "I", "K", "M", "O", "Q", "S"):
        label = ws[f"{col_letter}{SIDE_REF_DAYS_ROW}"].value
        assert isinstance(label, str) and len(label) == 3 and label.isupper(), (
            f"{col_letter}{SIDE_REF_DAYS_ROW} = {label!r}; expected a 3-letter day"
        )
    assert ws[f"C{SIDE_REF_DAYS_ROW}"].value == "Hi / Lo °F"


def test_reference_day_labels_align_with_the_modesto_header(tmp_path: Path):
    """Same column spans as row 11, so the band reads as part of one grid."""
    ws = _render_default(tmp_path)

    for col_letter in ("E", "G", "I", "K", "M", "O", "Q", "S"):
        assert ws[f"{col_letter}{SIDE_REF_DAYS_ROW}"].value == ws[f"{col_letter}11"].value, (
            f"{col_letter}: reference day label disagrees with the Modesto header"
        )


def test_banner_is_generic_and_marked_reference_only(tmp_path: Path):
    """The banner covers two cities now - it must not be titled after one of them."""
    ws = _render_default(tmp_path)

    banner = str(ws[f"C{SIDE_REF_BANNER_ROW}"].value)
    assert "SIDE REFERENCE TEMPS" in banner
    assert "PORTLAND" not in banner.upper()
    assert "PHOENIX" not in banner.upper()
    assert "not included in the Modesto consensus" in banner
    assert "Google Weather" in banner


def test_phoenix_row_is_shaded_differently_from_portland(tmp_path: Path):
    """Two adjacent reference rows in one colour would read as one source."""
    ws = _render_default(tmp_path)

    portland_label = _rgb(ws[f"C{PORTLAND_DATA_ROW}"])
    phoenix_label = _rgb(ws[f"C{PHOENIX_DATA_ROW}"])
    banner = _rgb(ws[f"C{SIDE_REF_BANNER_ROW}"])
    assert len({portland_label, phoenix_label, banner}) == 3, (
        f"label fills must all differ: portland={portland_label} phoenix={phoenix_label} banner={banner}"
    )

    for col_letter in DAY_COLUMNS:
        portland_fill = _rgb(ws[f"{col_letter}{PORTLAND_DATA_ROW}"])
        phoenix_fill = _rgb(ws[f"{col_letter}{PHOENIX_DATA_ROW}"])
        assert portland_fill != phoenix_fill, f"{col_letter}: rows share fill {portland_fill}"

    # The shared header is the same colour above both rows
    assert _rgb(ws[f"C{SIDE_REF_DAYS_ROW}"]) == banner


def test_reference_values_are_excluded_from_the_modesto_weighted_average(tmp_path: Path):
    """The consensus formula must still span only source rows 13-18."""
    ws = _render_default(tmp_path)

    for col_letter in ("E", "F", "G", "H"):
        formula = ws[f"{col_letter}19"].value
        assert isinstance(formula, str) and formula.startswith("=")
        assert f"{col_letter}13:{col_letter}18" in formula
        rest = formula.replace(f"{col_letter}13:{col_letter}18", "")
        for row in (PORTLAND_DATA_ROW, PHOENIX_DATA_ROW):
            assert f"{col_letter}{row}" not in rest


def test_missing_reference_data_blanks_both_rows_without_crashing(tmp_path: Path):
    ws = _render(tmp_path, None, None)

    assert ws[f"C{PORTLAND_DATA_ROW}"].value == "PORTLAND, OR"
    assert ws[f"C{PHOENIX_DATA_ROW}"].value == "PHOENIX, AZ"
    assert _row_values(ws, PORTLAND_DATA_ROW) == ["--"] * 16
    assert _row_values(ws, PHOENIX_DATA_ROW) == ["--"] * 16


def test_one_missing_city_does_not_blank_the_other(tmp_path: Path):
    _, _, portland_data, phoenix_data, _ = _inputs()
    ws = _render(tmp_path, None, phoenix_data)

    assert _row_values(ws, PORTLAND_DATA_ROW) == ["--"] * 16
    assert ws[f"E{PHOENIX_DATA_ROW}"].value == 105


def test_reference_gaps_render_as_placeholders_not_shifted_values(tmp_path: Path):
    """A short forecast must leave later columns blank, not slide left."""
    _, _, portland_data, phoenix_data, _ = _inputs()
    portland_data["daily"] = portland_data["daily"][:3]
    phoenix_data["daily"] = phoenix_data["daily"][:2]
    ws = _render(tmp_path, portland_data, phoenix_data)

    assert ws[f"E{PORTLAND_DATA_ROW}"].value == 78
    assert ws[f"I{PORTLAND_DATA_ROW}"].value == 80   # day index 2 high
    assert ws[f"K{PORTLAND_DATA_ROW}"].value == "--"  # day index 3 high, no data

    assert ws[f"E{PHOENIX_DATA_ROW}"].value == 105
    assert ws[f"G{PHOENIX_DATA_ROW}"].value == 106    # day index 1 high
    assert ws[f"I{PHOENIX_DATA_ROW}"].value == "--"   # day index 2 high, no data


def test_solar_block_shifted_below_the_reference_band(tmp_path: Path):
    ws = _render_default(tmp_path)

    assert "SOLAR FORECAST" in str(ws[f"D{SOLAR_TITLE_ROW}"].value)
    assert ws[f"D{SOLAR_HEADER_ROW}"].value == "DATE"
    assert ws[f"E{SOLAR_HEADER_ROW}"].value == "9AM"
    # One empty spacer row between the band and the solar title
    for c in "CDEFGH":
        assert ws[f"{c}{PHOENIX_DATA_ROW + 1}"].value in (None, "")
    # Solar data must not collide with the band above it
    for row in (PORTLAND_DATA_ROW, PHOENIX_DATA_ROW):
        assert ws[f"D{row}"].value in (None, "")


def test_legend_row_lands_after_the_seven_day_solar_block(tmp_path: Path):
    ws = _render_default(tmp_path)

    assert ws[f"E{SOLAR_LEGEND_ROW}"].value == "Tule Fog"
    assert ws[f"L{SOLAR_LEGEND_ROW}"].value == "Full Sun"
